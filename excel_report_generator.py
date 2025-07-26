import os
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

class ExcelReportGenerator:
    def __init__(self, workdir: str, template_path: str, csv_file: str):
        self.workdir = workdir
        self.template_path = template_path
        self.csv_file = csv_file

    def generate_excel_report(
        self, json_file: str, excel_file: str, project_uid: str, project_name: str,
        branch: str, commit: str, spoc: str, leak_count: int, assessment_type: str
    ):
        if not os.path.exists(self.template_path):
            print(f"ERROR: Excel template not found at {self.template_path}")
            return

        if not os.path.exists(json_file):
            print(f"Warning: {json_file} not found. Skipping Excel generation.")
            return

        try:
            workbook = load_workbook(self.template_path)
        except Exception as e:
            print(f"ERROR: Could not load Excel template: {e}")
            return

        with open(json_file, 'r') as f:
            findings = json.load(f)

        # --- Populate Executive Summary (Sheet 1) ---
        # This part needs to be customized based on your template's specific cells and variables.
        summary_sheet = workbook["Executive_Summary"]
        summary_sheet['D4'] = project_uid
        summary_sheet['D5'] = pd.Timestamp.now().strftime("%Y-%m-%d") # Scan Date
        summary_sheet['D6'] = "" # Close Date (initially empty)
        summary_sheet['D7'] = project_name
        summary_sheet['D8'] = branch
        summary_sheet['D9'] = commit
        summary_sheet['D10'] = spoc
        summary_sheet['D11'] = assessment_type # Assuming assessment_type is passed to this function

        # --- Populate Actual Findings (Sheet 2) ---
        findings_sheet = workbook["Actual_Findings"]

        # Define headers (adjust as per your desired order and names)
        headers = ["RuleID", "Description", "File", "Line", "Secret", "Severity", "Status"]
        #findings_sheet.append(headers)

        # Apply some basic styling to headers
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        for col_num, header_text in enumerate(headers, 1):
            cell = findings_sheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill

        # Get repo_base_url for hyperlinks
        repo_base_url = ""
        try:
            projects_df = pd.read_csv(os.path.join(self.workdir, self.csv_file))
            repo_base_url_series = projects_df.loc[projects_df["project_name"] == project_name, "git_url"]
            if not repo_base_url_series.empty:
                repo_base_url = repo_base_url_series.iloc[0].replace(".git", "")
        except FileNotFoundError:
            print(f"Warning: {self.csv_file} not found. Cannot determine repo_base_url for Excel.")

        # Add findings rows
        for finding in findings:
            rule_id = finding.get("RuleID", "N/A")
            description = finding.get("Description", "N/A")
            file = finding.get("File", "N/A").replace("/repo", "")
            start_line = finding.get("StartLine", "N/A")
            secret = finding.get("Secret", "N/A")
            author = finding.get("Author", "N/A")
            date = finding.get("Date", "N/A")[:10] if finding.get("Date") else "N/A"

            view_link = ""
            if commit != "Latest" and repo_base_url and file != "N/A" and start_line != "N/A":
                clean_file = file.lstrip('/')
                view_link = f"{repo_base_url}/-/blob/{commit}/{clean_file}#L{start_line}"

            row_data = [
                rule_id, description, file, start_line, secret, "", "Open" # Default status
            ]
            findings_sheet.append(row_data)

            # Add hyperlink to File column
            file_cell = findings_sheet.cell(row=findings_sheet.max_row, column=headers.index("File") + 1)
            if view_link:
                file_cell.hyperlink = view_link
                file_cell.style = "Hyperlink"
                file_cell.font = Font(color="0000FF", underline="single")

        # Add Data Validation for Status column
        status_col_letter = get_column_letter(headers.index("Status") + 1)
        dv = DataValidation(type="list", formula1='"Open,Closed"', allow_blank=True)
        # Ensure data validation range is valid by applying only when there are at least two rows
        if findings_sheet.max_row >= 2:
            dv.add(f'{status_col_letter}2:{status_col_letter}{findings_sheet.max_row}')
            findings_sheet.add_data_validation(dv)

        # Auto-fit columns (optional, can be slow for very large sheets)
        for col in findings_sheet.columns:
            max_length = 0
            column = col[0].column_letter # Get the column name
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.2
            findings_sheet.column_dimensions[column].width = adjusted_width

        try:
            workbook.save(excel_file)
            print(f"Generated Excel Report: {excel_file}")
        except Exception as e:
            print(f"ERROR: Could not save Excel report: {e}")
